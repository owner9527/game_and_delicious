"""把 Markdown 意见书转成 Word 文档（支持标题/正文/项目符号/表格/粗体，并正确设置中文字体）。
用法: python md2docx.py <input.md> <output.docx>
"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CJK = '微软雅黑'
LATIN = 'Segoe UI'


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = LATIN
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    rf.set(qn('w:ascii'), LATIN)
    rf.set(qn('w:hAnsi'), LATIN)
    rf.set(qn('w:eastAsia'), CJK)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_rich(par, text, size=10.5, base_bold=False):
    """处理 **粗体** 与 `代码` 两种行内标记。"""
    for chunk in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not chunk:
            continue
        if chunk.startswith('**') and chunk.endswith('**') and len(chunk) > 4:
            r = par.add_run(chunk[2:-2])
            set_run_font(r, size, True)
        elif chunk.startswith('`') and chunk.endswith('`') and len(chunk) > 2:
            r = par.add_run(chunk[1:-1])
            set_run_font(r, size - 0.5, False)
            r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        else:
            r = par.add_run(chunk)
            set_run_font(r, size, base_bold if base_bold else None)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def build(md_path, out_path):
    lines = open(md_path, encoding='utf-8').read().split('\n')
    doc = Document()

    # 页面与默认样式
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.4); sec.right_margin = Cm(2.4)

    normal = doc.styles['Normal']
    normal.font.name = LATIN
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), CJK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    for name, size in (('Heading 1', 17), ('Heading 2', 14), ('Heading 3', 12)):
        st = doc.styles[name]
        st.font.name = LATIN
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x36, 0x5D)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), CJK)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # 表格
        if line.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:\-|]+\|$', lines[i + 1].strip()):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not re.match(r'^[\s:\-|]+$', '|'.join(cells)):
                    rows.append(cells)
                i += 1
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=0, cols=ncol)
            t.style = 'Light Grid Accent 1'
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                cells = t.add_row().cells
                for ci in range(ncol):
                    txt = row[ci] if ci < len(row) else ''
                    p = cells[ci].paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    add_rich(p, txt, 9.5, base_bold=(ri == 0))
                    if ri == 0:
                        shade(cells[ci], 'DCE6F1')
            doc.add_paragraph()
            continue

        if not line.strip():
            i += 1; continue

        if re.match(r'^-{3,}$', line.strip()):
            i += 1; continue

        m = re.match(r'^(#{1,4})\s+(.*)$', line)
        if m:
            lvl = len(m.group(1))
            text = m.group(2).strip()
            if lvl == 1:
                p = doc.add_heading('', level=0)
                add_rich(p, text, 26)
            else:
                p = doc.add_heading('', level=min(3, lvl - 1))
                add_rich(p, text, {2: 17, 3: 14, 4: 12}.get(lvl, 12))
            i += 1; continue

        m = re.match(r'^[-*]\s+(.*)$', line)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            add_rich(p, m.group(1))
            i += 1; continue

        m = re.match(r'^\d+\.\s+(.*)$', line)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_rich(p, m.group(1))
            i += 1; continue

        p = doc.add_paragraph()
        add_rich(p, line.strip())
        i += 1

    # 页脚页码
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    r._element.addnext(fld)
    set_run_font(fp.add_run('  /  '), 9)
    r2 = fp.add_run()
    fld2 = OxmlElement('w:fldSimple')
    fld2.set(qn('w:instr'), 'NUMPAGES')
    r2._element.addnext(fld2)

    doc.save(out_path)
    print('saved', out_path)


if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2])
